# -*- coding: utf-8 -*-
"""
Basin Water Storage Degree (BWSD) calculation for the Fengshuwan catchment.

This script accompanies the manuscript:

    Basin Water Storage Degree: A Forward-Looking Water-Balance Indicator
    for Hydrological-State Warning in Small Debris-Flow Catchments

Purpose
-------
This script reproduces the core daily BWSD calculation using the processed
daily hydrological dataset for the Fengshuwan catchment.

The authoritative archived version of the processed daily hydrological dataset
is available on Zenodo:

    Wang, T. (2026). Processed Daily Hydrological Dataset for the Fengshuwan
    Catchment, Zhejiang, China, 2023–2026 (v1.0.0) [Data set]. Zenodo.
    https://doi.org/10.5281/zenodo.20068173

The script scope is limited to the core BWSD water-balance calculation,
forward-looking threshold construction, hydrological-state level assignment,
exceedance statistics, and summary-table export.

Input workbook
--------------
    data/fengshuwan_processed_daily_hydrology.xlsx

Input sheet
-----------
    Daily_Data

Required input columns
----------------------
    Date
    Precipitation_mm
    Evapotranspiration_mm
    Water_level_m
    Runoff_mm

Main equations
--------------
1. Apparent net water input:

       I_star[t] = P[t] - Q[t] - E[t]

2. Apparent basin storage increment:

       DeltaS[t0] = 0
       DeltaS[t] = max(0, DeltaS[t-1] + I_star[t])

3. Forward-looking critical-reference storage threshold:

       Scrit_minus[t] = max(S_min, S_init, S_stat_minus[t], S_nc_minus[t])

   where all threshold components at time t are computed only from information
   available before t.

4. Basin Water Storage Degree:

       BWSD[t] = DeltaS[t] / Scrit_minus[t]

Output workbook
---------------
    results/fengshuwan_bwsd_results.xlsx

Output sheets
-------------
    Daily_BWSD
    Exceedance
    Parameters
    Metadata
    Column_Dictionary

Run
---
    python main.py
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter


SCRIPT_NAME = Path(__file__).name

DATASET_TITLE = (
    "Processed Daily Hydrological Dataset for the Fengshuwan Catchment, "
    "Zhejiang, China, 2023–2026"
)

DATASET_VERSION = "v1.0.0"
DATASET_DOI = "10.5281/zenodo.20068173"
DATASET_URL = "https://doi.org/10.5281/zenodo.20068173"

DATASET_CITATION = (
    "Wang, T. (2026). Processed Daily Hydrological Dataset for the "
    "Fengshuwan Catchment, Zhejiang, China, 2023–2026 "
    "(v1.0.0) [Data set]. Zenodo. "
    "https://doi.org/10.5281/zenodo.20068173"
)


# =============================================================================
# 1. Parameters
# =============================================================================

@dataclass(frozen=True)
class BWSDParameters:
    """Fixed parameter set for the Fengshuwan daily BWSD case calculation."""

    catchment_name: str = "Fengshuwan catchment"
    catchment_area_km2: float = 0.8461

    # Archived input dataset information.
    dataset_title: str = DATASET_TITLE
    dataset_version: str = DATASET_VERSION
    dataset_doi: str = DATASET_DOI
    dataset_url: str = DATASET_URL
    dataset_citation: str = DATASET_CITATION

    # Low-storage reference date.
    t0: pd.Timestamp = pd.Timestamp("2023-10-12")

    # Initialization settings.
    hydrological_initialization_days: int = 90
    formal_initialization_days: int = 365

    # Threshold-construction parameters.
    p_stat: float = 0.98
    m0: float = 1.20
    m_nc: float = 1.20
    s_min_mm: float = 20.0

    # Management-oriented BWSD levels.
    theta_1: float = 0.70
    theta_2: float = 0.90

    # Minimum prior samples for threshold components.
    min_prior_for_stat: int = 30
    min_prior_for_nc: int = 30

    # Missing-flux policy.
    # "carry": keep previous DeltaS if P, E, or Q is missing.
    # "nan": set DeltaS to NaN on missing-flux days.
    missing_flux_policy: str = "carry"

    @property
    def raw_diagnostic_start(self) -> pd.Timestamp:
        """Start date after hydrological initialization."""
        return self.t0 + pd.Timedelta(days=self.hydrological_initialization_days)

    @property
    def formal_start(self) -> pd.Timestamp:
        """Start date of the formal forward-looking evaluation period."""
        return self.t0 + pd.Timedelta(days=self.formal_initialization_days)


# =============================================================================
# 2. Helper functions
# =============================================================================

def finite_values(values) -> np.ndarray:
    """Return finite numeric values from an array-like object."""
    arr = np.asarray(values, dtype=float)
    return arr[np.isfinite(arr)]


def safe_quantile(values, q: float) -> float:
    """Return quantile from finite values, or NaN if no finite value exists."""
    vals = finite_values(values)
    if vals.size == 0:
        return np.nan
    return float(np.quantile(vals, q))


def safe_max(values) -> float:
    """Return maximum from finite values, or NaN if no finite value exists."""
    vals = finite_values(values)
    if vals.size == 0:
        return np.nan
    return float(np.max(vals))


def warning_level(bwsd: float, theta_1: float, theta_2: float) -> float:
    """
    Convert BWSD into a management-oriented hydrological level.

    Levels
    ------
    0: normal hydrological state
    1: enhanced storage state
    2: high-storage sensitive state
    3: critical-reference hydrological state
    """
    if not np.isfinite(bwsd):
        return np.nan

    if bwsd < theta_1:
        return 0

    if bwsd < theta_2:
        return 1

    if bwsd < 1.0:
        return 2

    return 3


def evaluation_phase(date: pd.Timestamp, params: BWSDParameters) -> str:
    """Assign the evaluation phase for a given date."""
    if date < params.t0:
        return "Pre_reference"

    if date < params.raw_diagnostic_start:
        return "Hydrological_initialization"

    if date < params.formal_start:
        return "Threshold_initialization"

    return "Formal"


def max_value_and_date(df: pd.DataFrame, value_column: str) -> tuple[float, str]:
    """Return the maximum value and corresponding date string for a column."""
    if value_column not in df.columns:
        return np.nan, ""

    s = pd.to_numeric(df[value_column], errors="coerce")

    if not s.notna().any():
        return np.nan, ""

    idx = int(s.idxmax())
    value = float(s.loc[idx])
    date_string = pd.Timestamp(df.loc[idx, "Date"]).strftime("%Y-%m-%d")

    return value, date_string


def autofit_worksheet(ws, min_width: int = 10, max_width: int = 48) -> None:
    """Auto-fit Excel worksheet column widths."""
    widths = {}

    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row, start=1):
            text = "" if value is None else str(value)
            widths[idx] = max(
                widths.get(idx, min_width),
                min(len(text) + 3, max_width),
            )

    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(
            min_width,
            min(width, max_width),
        )


# =============================================================================
# 3. Input loading
# =============================================================================

def load_daily_hydrology(input_xlsx: Path, sheet_name: str) -> pd.DataFrame:
    """
    Load the processed daily hydrological dataset.

    Parameters
    ----------
    input_xlsx : pathlib.Path
        Path to the input workbook.
    sheet_name : str
        Name of the input sheet.

    Returns
    -------
    pandas.DataFrame
        Cleaned daily hydrological data.
    """
    input_xlsx = Path(input_xlsx)

    if not input_xlsx.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_xlsx}")

    df = pd.read_excel(input_xlsx, sheet_name=sheet_name)

    required_columns = [
        "Date",
        "Precipitation_mm",
        "Evapotranspiration_mm",
        "Water_level_m",
        "Runoff_mm",
    ]

    missing_columns = [col for col in required_columns if col not in df.columns]

    if missing_columns:
        raise KeyError(f"Missing required input columns: {missing_columns}")

    out = pd.DataFrame()
    out["Date"] = pd.to_datetime(df["Date"], errors="coerce").dt.normalize()
    out["Precipitation_mm"] = pd.to_numeric(df["Precipitation_mm"], errors="coerce")
    out["Evapotranspiration_mm"] = pd.to_numeric(
        df["Evapotranspiration_mm"],
        errors="coerce",
    )
    out["Water_level_m"] = pd.to_numeric(df["Water_level_m"], errors="coerce")
    out["Runoff_mm"] = pd.to_numeric(df["Runoff_mm"], errors="coerce")

    out = (
        out
        .dropna(subset=["Date"])
        .sort_values("Date")
        .reset_index(drop=True)
    )

    duplicated_dates = out["Date"].duplicated(keep=False)
    if duplicated_dates.any():
        duplicated_list = (
            out.loc[duplicated_dates, "Date"]
            .dt.strftime("%Y-%m-%d")
            .unique()
            .tolist()
        )
        raise ValueError(f"Duplicated dates were found: {duplicated_list}")

    if out.empty:
        raise ValueError("The input sheet contains no valid daily records.")

    return out


# =============================================================================
# 4. BWSD calculation
# =============================================================================

def calculate_bwsd(
    hydrology: pd.DataFrame,
    params: BWSDParameters,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """
    Calculate daily BWSD using a strictly forward-looking threshold rule.

    Parameters
    ----------
    hydrology : pandas.DataFrame
        Processed daily hydrological data.
    params : BWSDParameters
        Fixed BWSD parameter set.

    Returns
    -------
    tuple
        daily_bwsd, exceedance_table, metadata_table
    """
    df = pd.DataFrame()

    df["Date"] = hydrology["Date"]
    df["Evaluation_phase"] = df["Date"].apply(lambda d: evaluation_phase(d, params))
    df["Is_formal"] = df["Date"] >= params.formal_start

    # Processed hydrological inputs.
    df["Precipitation_mm"] = hydrology["Precipitation_mm"]
    df["Evapotranspiration_mm"] = hydrology["Evapotranspiration_mm"]
    df["Water_level_m"] = hydrology["Water_level_m"]
    df["Runoff_mm"] = hydrology["Runoff_mm"]

    # Short aliases used in water-balance calculation.
    df["P_mm"] = df["Precipitation_mm"]
    df["E_mm"] = df["Evapotranspiration_mm"]
    df["Q_mm"] = df["Runoff_mm"]

    df["Flux_complete"] = (
        df["P_mm"].notna()
        & df["E_mm"].notna()
        & df["Q_mm"].notna()
    )

    # No debris-flow event was documented in this monitoring record.
    # No abnormal samples are released in the simplified public dataset.
    df["Debris_flow_event"] = False
    df["Abnormal_sample"] = False

    df["Valid_for_threshold"] = (
        df["Flux_complete"]
        & (~df["Debris_flow_event"])
        & (~df["Abnormal_sample"])
    )

    # In the present no-event case, prior complete and valid samples are treated
    # as non-event samples after they have occurred. The forward-looking rule is
    # enforced below by using only prior samples j < t.
    df["Non_event_threshold_sample"] = df["Valid_for_threshold"]

    # Apparent net water input.
    df["I_star_mm"] = np.where(
        df["Flux_complete"],
        df["P_mm"] - df["Q_mm"] - df["E_mm"],
        np.nan,
    )

    # -------------------------------------------------------------------------
    # 4.1 Apparent basin storage increment
    # -------------------------------------------------------------------------
    dates = df["Date"]

    if not dates.eq(params.t0).any():
        raise ValueError(
            f"T0 was not found in the input date sequence: {params.t0.date()}"
        )

    t0_idx = int(df.index[dates.eq(params.t0)][0])
    n = len(df)

    delta_s = np.full(n, np.nan, dtype=float)

    for i in range(n):
        if i < t0_idx:
            delta_s[i] = np.nan
            continue

        if i == t0_idx:
            delta_s[i] = 0.0
            continue

        previous_delta_s = delta_s[i - 1]

        if not np.isfinite(previous_delta_s):
            previous_delta_s = 0.0

        i_star = df.loc[i, "I_star_mm"]

        if np.isfinite(i_star):
            delta_s[i] = max(0.0, previous_delta_s + float(i_star))
        else:
            policy = params.missing_flux_policy.lower()

            if policy == "carry":
                delta_s[i] = previous_delta_s
            elif policy == "nan":
                delta_s[i] = np.nan
            else:
                raise ValueError(
                    "Unsupported missing_flux_policy. Use 'carry' or 'nan'."
                )

    df["DeltaS_mm"] = delta_s

    # -------------------------------------------------------------------------
    # 4.2 Initialization lower bound
    # -------------------------------------------------------------------------
    init_mask = (
        (df["Date"] >= params.t0)
        & (df["Date"] < params.formal_start)
        & df["Valid_for_threshold"]
        & df["DeltaS_mm"].notna()
    )

    if not init_mask.any():
        raise ValueError("No valid initialization samples were found.")

    initialization_sample_days = int(init_mask.sum())
    initialization_max_delta_s = float(df.loc[init_mask, "DeltaS_mm"].max())
    s_init = max(params.s_min_mm, params.m0 * initialization_max_delta_s)

    # -------------------------------------------------------------------------
    # 4.3 Forward-looking critical-reference threshold
    # -------------------------------------------------------------------------
    s_stat_minus = np.full(n, np.nan, dtype=float)
    s_nc_minus = np.full(n, np.nan, dtype=float)
    s_crit_minus = np.full(n, np.nan, dtype=float)
    threshold_control = []

    index_array = np.arange(n)
    valid_array = df["Valid_for_threshold"].to_numpy(dtype=bool)
    non_event_array = df["Non_event_threshold_sample"].to_numpy(dtype=bool)
    delta_array = df["DeltaS_mm"].to_numpy(dtype=float)

    for i in range(n):
        # Strictly prior information only: j < i.
        prior_mask = (
            (index_array < i)
            & valid_array
            & np.isfinite(delta_array)
        )

        prior_values = delta_array[prior_mask]

        if finite_values(prior_values).size >= params.min_prior_for_stat:
            s_stat_minus[i] = safe_quantile(prior_values, params.p_stat)

        prior_non_event_mask = prior_mask & non_event_array
        prior_non_event_values = delta_array[prior_non_event_mask]

        if finite_values(prior_non_event_values).size >= params.min_prior_for_nc:
            s_nc_minus[i] = params.m_nc * safe_max(prior_non_event_values)

        components = {
            "S_min": params.s_min_mm,
            "S_init": s_init,
        }

        if np.isfinite(s_stat_minus[i]):
            components["S_stat"] = s_stat_minus[i]

        if np.isfinite(s_nc_minus[i]):
            components["S_nc"] = s_nc_minus[i]

        controlling_component = max(components, key=components.get)
        s_crit_minus[i] = float(components[controlling_component])
        threshold_control.append(controlling_component)

    df["S_stat_minus_mm"] = s_stat_minus
    df["S_nc_minus_mm"] = s_nc_minus
    df["S_init_mm"] = s_init
    df["S_crit_minus_mm"] = s_crit_minus
    df["Threshold_control"] = threshold_control

    # -------------------------------------------------------------------------
    # 4.4 BWSD and warning levels
    # -------------------------------------------------------------------------
    df["BWSD_all"] = df["DeltaS_mm"] / df["S_crit_minus_mm"]

    # Formal BWSD is retained only after the formal evaluation start.
    df["BWSD_formal"] = np.where(
        df["Date"] >= params.formal_start,
        df["BWSD_all"],
        np.nan,
    )

    df["BWSD_plot"] = df["BWSD_all"].clip(lower=0.0, upper=1.20)
    df["BWSD_0_1"] = df["BWSD_all"].clip(lower=0.0, upper=1.00)

    df["BWSD_level_all"] = df["BWSD_all"].apply(
        lambda x: warning_level(x, params.theta_1, params.theta_2)
    )

    df["BWSD_level_formal"] = df["BWSD_formal"].apply(
        lambda x: warning_level(x, params.theta_1, params.theta_2)
    )

    df["Theta_1"] = params.theta_1
    df["Theta_2"] = params.theta_2
    df["Critical_reference_level"] = 1.0

    # -------------------------------------------------------------------------
    # 4.5 Exceedance table
    # -------------------------------------------------------------------------
    exceedance_table = build_exceedance_table(df)

    # -------------------------------------------------------------------------
    # 4.6 Metadata
    # -------------------------------------------------------------------------
    metadata_table = build_metadata_table(
        df=df,
        params=params,
        initialization_sample_days=initialization_sample_days,
        initialization_max_delta_s=initialization_max_delta_s,
        s_init=s_init,
    )

    return df, exceedance_table, metadata_table


# =============================================================================
# 5. Summary tables
# =============================================================================

def build_exceedance_table(df: pd.DataFrame) -> pd.DataFrame:
    """Build exceedance statistics for selected BWSD levels."""
    thresholds = [0.50, 0.60, 0.70, 0.80, 0.85, 0.90, 0.95, 1.00]

    rows = []

    formal_total_days = int(df["BWSD_formal"].notna().sum())

    for threshold in thresholds:
        all_days = int((df["BWSD_all"] >= threshold).sum())
        formal_days = int((df["BWSD_formal"] >= threshold).sum())

        if formal_total_days > 0:
            formal_percent = 100.0 * formal_days / formal_total_days
        else:
            formal_percent = np.nan

        rows.append(
            {
                "Threshold": threshold,
                "All_period_exceedance_days": all_days,
                "Formal_exceedance_days": formal_days,
                "Formal_total_days": formal_total_days,
                "Formal_exceedance_percent": round(formal_percent, 3),
            }
        )

    return pd.DataFrame(rows)


def build_metadata_table(
    df: pd.DataFrame,
    params: BWSDParameters,
    initialization_sample_days: int,
    initialization_max_delta_s: float,
    s_init: float,
) -> pd.DataFrame:
    """Build metadata and key result summary table."""
    cumulative_p = float(df["P_mm"].sum(skipna=True))
    cumulative_e = float(df["E_mm"].sum(skipna=True))
    cumulative_q = float(df["Q_mm"].sum(skipna=True))
    apparent_residual = cumulative_p - cumulative_e - cumulative_q

    max_delta_s, date_max_delta_s = max_value_and_date(df, "DeltaS_mm")
    max_bwsd_all, date_max_bwsd_all = max_value_and_date(df, "BWSD_all")
    max_bwsd_formal, date_max_bwsd_formal = max_value_and_date(df, "BWSD_formal")

    formal_days = int((df["Date"] >= params.formal_start).sum())
    formal_valid_days = int(df["BWSD_formal"].notna().sum())

    formal_ge_070 = int((df["BWSD_formal"] >= 0.70).sum())
    formal_ge_080 = int((df["BWSD_formal"] >= 0.80).sum())
    formal_ge_085 = int((df["BWSD_formal"] >= 0.85).sum())
    formal_ge_090 = int((df["BWSD_formal"] >= 0.90).sum())
    formal_ge_095 = int((df["BWSD_formal"] >= 0.95).sum())
    formal_ge_100 = int((df["BWSD_formal"] >= 1.00).sum())

    rows = [
        ("Script", SCRIPT_NAME),
        ("Generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Catchment_name", params.catchment_name),
        ("Catchment_area_km2", params.catchment_area_km2),
        ("Input_dataset_title", params.dataset_title),
        ("Input_dataset_version", params.dataset_version),
        ("Input_dataset_DOI", params.dataset_doi),
        ("Input_dataset_URL", params.dataset_url),
        ("Input_dataset_citation", params.dataset_citation),
        ("Date_range", f"{df['Date'].min().date()} to {df['Date'].max().date()}"),
        ("Total_days", int(len(df))),
        ("T0", params.t0.strftime("%Y-%m-%d")),
        ("Raw_diagnostic_start", params.raw_diagnostic_start.strftime("%Y-%m-%d")),
        ("Formal_evaluation_start", params.formal_start.strftime("%Y-%m-%d")),
        ("Hydrological_initialization_days", params.hydrological_initialization_days),
        ("Formal_initialization_days", params.formal_initialization_days),
        ("Initialization_sample_days", initialization_sample_days),
        ("P_STAT", params.p_stat),
        ("M0", params.m0),
        ("M_NC", params.m_nc),
        ("S_MIN_mm", params.s_min_mm),
        ("Theta_1", params.theta_1),
        ("Theta_2", params.theta_2),
        ("Minimum_prior_samples_for_stat", params.min_prior_for_stat),
        ("Minimum_prior_samples_for_non_event", params.min_prior_for_nc),
        ("Missing_flux_policy", params.missing_flux_policy),
        ("Cumulative_precipitation_mm", round(cumulative_p, 6)),
        ("Cumulative_evapotranspiration_mm", round(cumulative_e, 6)),
        ("Cumulative_runoff_mm", round(cumulative_q, 6)),
        ("Apparent_water_balance_residual_mm", round(apparent_residual, 6)),
        ("Initialization_max_DeltaS_mm", round(initialization_max_delta_s, 6)),
        ("S_init_mm", round(s_init, 6)),
        ("Max_DeltaS_mm", round(max_delta_s, 6)),
        ("Date_of_max_DeltaS", date_max_delta_s),
        ("Max_BWSD_all", round(max_bwsd_all, 6)),
        ("Date_of_max_BWSD_all", date_max_bwsd_all),
        ("Max_BWSD_formal", round(max_bwsd_formal, 6)),
        ("Date_of_max_BWSD_formal", date_max_bwsd_formal),
        ("Formal_days", formal_days),
        ("Formal_valid_BWSD_days", formal_valid_days),
        ("Formal_BWSD_ge_0.70_days", formal_ge_070),
        ("Formal_BWSD_ge_0.80_days", formal_ge_080),
        ("Formal_BWSD_ge_0.85_days", formal_ge_085),
        ("Formal_BWSD_ge_0.90_days", formal_ge_090),
        ("Formal_BWSD_ge_0.95_days", formal_ge_095),
        ("Formal_BWSD_ge_1.00_days", formal_ge_100),
        ("Documented_debris_flow_events", 0),
        (
            "Forward_looking_rule",
            "For each date t, S_crit_minus[t] is calculated only from prior samples j < t.",
        ),
        (
            "Code_scope",
            "This script performs the core BWSD calculation and summary export.",
        ),
    ]

    return pd.DataFrame(rows, columns=["Item", "Value"])


def build_parameters_table(params: BWSDParameters) -> pd.DataFrame:
    """Return the fixed parameter table."""
    rows = [
        (
            "Catchment_area_km2",
            params.catchment_area_km2,
            "km2",
            "Catchment drainage area.",
        ),
        (
            "T0",
            params.t0.strftime("%Y-%m-%d"),
            "date",
            "Low-storage reference date.",
        ),
        (
            "Hydrological_initialization_days",
            params.hydrological_initialization_days,
            "day",
            "Diagnostic hydrological initialization length.",
        ),
        (
            "Formal_initialization_days",
            params.formal_initialization_days,
            "day",
            "Threshold initialization and learning length before formal evaluation.",
        ),
        (
            "P_STAT",
            params.p_stat,
            "-",
            "Quantile level for statistical prior threshold.",
        ),
        (
            "M0",
            params.m0,
            "-",
            "Safety coefficient for initialization lower bound.",
        ),
        (
            "M_NC",
            params.m_nc,
            "-",
            "Safety coefficient for non-event threshold constraint.",
        ),
        (
            "S_MIN_mm",
            params.s_min_mm,
            "mm",
            "Minimum positive critical-reference threshold.",
        ),
        (
            "Theta_1",
            params.theta_1,
            "-",
            "Management threshold for enhanced storage state.",
        ),
        (
            "Theta_2",
            params.theta_2,
            "-",
            "Management threshold for high-storage sensitive state.",
        ),
        (
            "Min_prior_for_stat",
            params.min_prior_for_stat,
            "sample",
            "Minimum prior samples required for statistical threshold.",
        ),
        (
            "Min_prior_for_nc",
            params.min_prior_for_nc,
            "sample",
            "Minimum prior non-event samples required for non-event threshold constraint.",
        ),
        (
            "Missing_flux_policy",
            params.missing_flux_policy,
            "-",
            "Storage-update rule when one or more flux components are missing.",
        ),
    ]

    return pd.DataFrame(rows, columns=["Parameter", "Value", "Unit", "Description"])


def build_column_dictionary() -> pd.DataFrame:
    """Return the output column dictionary."""
    rows = [
        ("Date", "Date of the daily record."),
        (
            "Evaluation_phase",
            "Evaluation phase assigned according to the reference and formal-evaluation dates.",
        ),
        (
            "Is_formal",
            "Whether the date belongs to the formal forward-looking evaluation period.",
        ),
        ("Precipitation_mm", "Processed daily precipitation depth."),
        ("Evapotranspiration_mm", "Processed daily evapotranspiration-loss depth."),
        ("Water_level_m", "Processed final outlet water level."),
        ("Runoff_mm", "Processed daily outlet runoff depth."),
        ("P_mm", "Precipitation alias used in the water-balance calculation."),
        ("E_mm", "Evapotranspiration alias used in the water-balance calculation."),
        ("Q_mm", "Runoff alias used in the water-balance calculation."),
        ("Flux_complete", "Whether P, E, and Q are all available."),
        (
            "Debris_flow_event",
            "Confirmed debris-flow event flag. No event was documented in this case.",
        ),
        (
            "Abnormal_sample",
            "Abnormal-sample flag. No abnormal sample is released in the simplified public dataset.",
        ),
        (
            "Valid_for_threshold",
            "Whether the sample is eligible for threshold construction after occurrence.",
        ),
        (
            "Non_event_threshold_sample",
            "Non-event sample flag used by the forward-looking non-event threshold constraint.",
        ),
        ("I_star_mm", "Apparent net water input, calculated as P - Q - E."),
        ("DeltaS_mm", "Apparent basin storage increment relative to T0."),
        ("S_stat_minus_mm", "Forward-looking statistical prior threshold."),
        ("S_nc_minus_mm", "Forward-looking non-event threshold constraint."),
        ("S_init_mm", "Initialization lower-bound threshold."),
        (
            "S_crit_minus_mm",
            "Final forward-looking critical-reference storage threshold.",
        ),
        (
            "Threshold_control",
            "Threshold component controlling the final S_crit_minus value.",
        ),
        ("BWSD_all", "BWSD values for all dates after T0; pre-formal values are diagnostic."),
        ("BWSD_formal", "BWSD values retained only for the formal evaluation period."),
        ("BWSD_plot", "BWSD_all clipped to 1.20 for optional plotting convenience."),
        ("BWSD_0_1", "BWSD_all clipped to 1.00 for optional classification convenience."),
        ("BWSD_level_all", "BWSD-based hydrological level for all dates."),
        ("BWSD_level_formal", "BWSD-based hydrological level for formal evaluation dates."),
        ("Theta_1", "Management threshold for enhanced storage state."),
        ("Theta_2", "Management threshold for high-storage sensitive state."),
        ("Critical_reference_level", "Critical-reference BWSD level, equal to 1.0."),
    ]

    return pd.DataFrame(rows, columns=["Column", "Description"])


# =============================================================================
# 6. Output
# =============================================================================

def write_results(
    daily_bwsd: pd.DataFrame,
    exceedance_table: pd.DataFrame,
    parameters_table: pd.DataFrame,
    metadata_table: pd.DataFrame,
    column_dictionary: pd.DataFrame,
    output_xlsx: Path,
) -> None:
    """Write BWSD outputs to an Excel workbook."""
    output_xlsx = Path(output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    daily_to_excel = daily_bwsd.copy()
    daily_to_excel["Date"] = pd.to_datetime(daily_to_excel["Date"]).dt.date

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        daily_to_excel.to_excel(writer, sheet_name="Daily_BWSD", index=False)
        exceedance_table.to_excel(writer, sheet_name="Exceedance", index=False)
        parameters_table.to_excel(writer, sheet_name="Parameters", index=False)
        metadata_table.to_excel(writer, sheet_name="Metadata", index=False)
        column_dictionary.to_excel(writer, sheet_name="Column_Dictionary", index=False)

        for worksheet in writer.sheets.values():
            autofit_worksheet(worksheet)


def print_key_results(metadata_table: pd.DataFrame, output_xlsx: Path) -> None:
    """Print key results to console."""
    lookup = dict(zip(metadata_table["Item"], metadata_table["Value"]))

    print("=" * 88)
    print("BWSD calculation completed")
    print("-" * 88)
    print(f"Output workbook                       : {output_xlsx}")
    print(f"Input dataset DOI                     : {lookup.get('Input_dataset_DOI')}")
    print(f"Date range                            : {lookup.get('Date_range')}")
    print(f"T0                                    : {lookup.get('T0')}")
    print(f"Formal evaluation start               : {lookup.get('Formal_evaluation_start')}")
    print(f"Initialization max DeltaS             : {lookup.get('Initialization_max_DeltaS_mm')} mm")
    print(f"S_init                                : {lookup.get('S_init_mm')} mm")
    print(f"Max DeltaS                            : {lookup.get('Max_DeltaS_mm')} mm")
    print(f"Date of max DeltaS                    : {lookup.get('Date_of_max_DeltaS')}")
    print(f"Max formal BWSD                       : {lookup.get('Max_BWSD_formal')}")
    print(f"Date of max formal BWSD               : {lookup.get('Date_of_max_BWSD_formal')}")
    print(f"Formal BWSD >= 0.70 days              : {lookup.get('Formal_BWSD_ge_0.70_days')}")
    print(f"Formal BWSD >= 0.90 days              : {lookup.get('Formal_BWSD_ge_0.90_days')}")
    print(f"Formal BWSD >= 1.00 days              : {lookup.get('Formal_BWSD_ge_1.00_days')}")
    print("=" * 88)


# =============================================================================
# 7. Main entry
# =============================================================================

def main() -> None:
    """Command-line entry point."""
    parser = argparse.ArgumentParser(
        description="Calculate Basin Water Storage Degree for the Fengshuwan catchment."
    )

    parser.add_argument(
        "--input",
        default=Path("data/fengshuwan_processed_daily_hydrology.xlsx"),
        type=Path,
        help="Path to the processed daily hydrological workbook.",
    )

    parser.add_argument(
        "--sheet",
        default="Daily_Data",
        help="Input sheet name.",
    )

    parser.add_argument(
        "--output",
        default=Path("results/fengshuwan_bwsd_results.xlsx"),
        type=Path,
        help="Path to the output BWSD result workbook.",
    )

    # parse_known_args makes the script more tolerant when executed from
    # interactive environments that may inject extra arguments.
    args, _ = parser.parse_known_args()

    params = BWSDParameters()

    hydrology = load_daily_hydrology(
        input_xlsx=args.input,
        sheet_name=args.sheet,
    )

    daily_bwsd, exceedance_table, metadata_table = calculate_bwsd(
        hydrology=hydrology,
        params=params,
    )

    parameters_table = build_parameters_table(params)
    column_dictionary = build_column_dictionary()

    write_results(
        daily_bwsd=daily_bwsd,
        exceedance_table=exceedance_table,
        parameters_table=parameters_table,
        metadata_table=metadata_table,
        column_dictionary=column_dictionary,
        output_xlsx=args.output,
    )

    print_key_results(metadata_table, args.output)


if __name__ == "__main__":
    main()