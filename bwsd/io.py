"""CSV/XLSX input, scientific validation, and workbook output."""

from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter
from .parameters import BWSDParameters
from .validation import validate_daily_hydrology


def autofit_worksheet(ws, min_width: int = 10, max_width: int = 48) -> None:
    """Auto-fit Excel worksheet column widths."""
    widths = {}

    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row, start=1):
            text = "" if value is None else str(value)
            widths[idx] = max(
                widths.get(idx, min_width),
                min(len(text) + 3, max_width),
            )

    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(
            min_width,
            min(width, max_width),
        )


def load_daily_hydrology(input_xlsx: Path, sheet_name: str) -> pd.DataFrame:
    """
    Load the processed daily hydrological dataset.

    Parameters
    ----------
    input_xlsx : pathlib.Path
        Path to the input workbook.
    sheet_name : str
        Name of the input sheet.

    Returns
    -------
    pandas.DataFrame
        Cleaned daily hydrological data.
    """
    input_xlsx = Path(input_xlsx)

    if not input_xlsx.exists():
        raise FileNotFoundError(f"Input workbook not found: {input_xlsx}")

    df = (
        pd.read_csv(input_xlsx)
        if input_xlsx.suffix.lower() == ".csv"
        else pd.read_excel(input_xlsx, sheet_name=sheet_name)
    )

    required_columns = [
        "Date",
        "Precipitation_mm",
        "Evapotranspiration_mm",
        "Water_level_m",
        "Runoff_mm",
    ]

    missing_columns = [col for col in required_columns if col not in df.columns]

    if missing_columns:
        raise KeyError(f"Missing required input columns: {missing_columns}")

    out = pd.DataFrame()
    out["Date"] = pd.to_datetime(df["Date"], errors="coerce").dt.normalize()
    if out["Date"].isna().any():
        raise ValueError("Invalid or missing observation date")
    out["Precipitation_mm"] = pd.to_numeric(df["Precipitation_mm"], errors="raise")
    out["Evapotranspiration_mm"] = pd.to_numeric(
        df["Evapotranspiration_mm"],
        errors="raise",
    )
    out["Water_level_m"] = pd.to_numeric(df["Water_level_m"], errors="raise")
    out["Runoff_mm"] = pd.to_numeric(df["Runoff_mm"], errors="raise")

    out = out.dropna(subset=["Date"]).sort_values("Date").reset_index(drop=True)

    duplicated_dates = out["Date"].duplicated(keep=False)
    if duplicated_dates.any():
        duplicated_list = (
            out.loc[duplicated_dates, "Date"].dt.strftime("%Y-%m-%d").unique().tolist()
        )
        raise ValueError(f"Duplicated dates were found: {duplicated_list}")

    if out.empty:
        raise ValueError("The input sheet contains no valid daily records.")

    for col in [
        "Abnormal_sample",
        "Debris_flow_event",
        "Noncritical_confirmed",
        "Confirmed_on",
    ]:
        if col in df.columns:
            out = out.merge(
                df[["Date", col]].assign(
                    Date=pd.to_datetime(df["Date"]).dt.normalize()
                ),
                on="Date",
                validate="one_to_one",
            )
    return validate_daily_hydrology(out)


def build_parameters_table(params: BWSDParameters) -> pd.DataFrame:
    """Return the fixed parameter table."""
    rows = [
        (
            "Catchment_area_km2",
            params.catchment_area_km2,
            "km2",
            "Catchment drainage area.",
        ),
        (
            "T0",
            params.t0.strftime("%Y-%m-%d"),
            "date",
            "Low-storage reference date.",
        ),
        (
            "Hydrological_initialization_days",
            params.hydrological_initialization_days,
            "day",
            "Diagnostic hydrological initialization length.",
        ),
        (
            "Formal_initialization_days",
            params.formal_initialization_days,
            "day",
            "Threshold initialization and learning length before formal evaluation.",
        ),
        (
            "P_STAT",
            params.p_stat,
            "-",
            "Quantile level for statistical prior threshold.",
        ),
        (
            "M0",
            params.m0,
            "-",
            "Safety coefficient for initialization lower bound.",
        ),
        (
            "M_NC",
            params.m_nc,
            "-",
            "Safety coefficient for non-event threshold constraint.",
        ),
        (
            "S_MIN_mm",
            params.s_min_mm,
            "mm",
            "Minimum positive critical-reference threshold.",
        ),
        (
            "Theta_1",
            params.theta_1,
            "-",
            "Management threshold for enhanced storage state.",
        ),
        (
            "Theta_2",
            params.theta_2,
            "-",
            "Management threshold for high-storage sensitive state.",
        ),
        (
            "Min_prior_for_stat",
            params.min_prior_for_stat,
            "sample",
            "Minimum prior samples required for statistical threshold.",
        ),
        (
            "Min_prior_for_nc",
            params.min_prior_for_nc,
            "sample",
            "Minimum qualified, independently confirmed high-storage responses for the field constraint.",
        ),
        (
            "Missing_flux_policy",
            params.missing_flux_policy,
            "-",
            "Storage-update rule when one or more flux components are missing.",
        ),
    ]

    return pd.DataFrame(rows, columns=["Parameter", "Value", "Unit", "Description"])


def build_column_dictionary() -> pd.DataFrame:
    """Return the output column dictionary."""
    rows = [
        ("Date", "Date of the daily record."),
        (
            "Evaluation_phase",
            "Evaluation phase assigned according to the reference and formal-evaluation dates.",
        ),
        (
            "Is_formal",
            "Whether the date belongs to the formal forward-looking evaluation period.",
        ),
        ("Precipitation_mm", "Processed daily precipitation depth."),
        (
            "Evapotranspiration_mm",
            "Evaporation-based atmospheric water-loss proxy; not measured catchment ET.",
        ),
        ("Water_level_m", "Processed final outlet water level."),
        ("Runoff_mm", "Processed daily outlet runoff depth."),
        ("P_mm", "Precipitation alias used in the water-balance calculation."),
        (
            "E_mm",
            "Evaporation-based atmospheric water-loss proxy; the archived column name is retained.",
        ),
        ("Q_mm", "Runoff alias used in the water-balance calculation."),
        ("Flux_complete", "Whether P, E, and Q are all available."),
        (
            "Debris_flow_event",
            "Confirmed debris-flow event flag. No event was documented in this case.",
        ),
        (
            "Abnormal_sample",
            "Abnormal-sample flag. No abnormal sample is released in the simplified public dataset.",
        ),
        (
            "Valid_for_threshold",
            "Whether the sample is eligible for threshold construction after occurrence.",
        ),
        (
            "Non_event_threshold_sample",
            "Independently confirmed high-storage sample qualified using its frozen BWSD; activates only after confirmation.",
        ),
        (
            "Noncritical_confirmed",
            "Independent non-critical confirmation flag; case defaults are explicit replay assumptions.",
        ),
        ("Confirmed_on", "Confirmation date; evidence activates only after this date."),
        ("I_star_mm", "Apparent net water input, calculated as P - Q - E."),
        ("DeltaS_mm", "Apparent basin storage increment relative to T0."),
        ("S_stat_minus_mm", "Forward-looking statistical prior threshold."),
        ("S_nc_minus_mm", "Forward-looking non-event threshold constraint."),
        ("S_init_mm", "Initialization lower-bound threshold."),
        (
            "S_crit_minus_mm",
            "Final forward-looking critical-reference storage threshold.",
        ),
        (
            "Threshold_control",
            "Threshold component controlling the final S_crit_minus value.",
        ),
        (
            "Preformal_retrospective_diagnostic",
            "Pre-formal ratios use the fixed initialization bound and are retrospective diagnostics.",
        ),
        (
            "BWSD_all",
            "BWSD values for all dates after T0; pre-formal values are diagnostic.",
        ),
        ("BWSD_formal", "BWSD values retained only for the formal evaluation period."),
        ("BWSD_plot", "BWSD_all clipped to 1.20 for optional plotting convenience."),
        (
            "BWSD_0_1",
            "BWSD_all clipped to 1.00 for optional classification convenience.",
        ),
        ("BWSD_level_all", "BWSD-based hydrological level for all dates."),
        (
            "BWSD_level_formal",
            "BWSD-based hydrological level for formal evaluation dates.",
        ),
        ("Theta_1", "Management threshold for enhanced storage state."),
        ("Theta_2", "Management threshold for high-storage sensitive state."),
        ("Critical_reference_level", "Critical-reference BWSD level, equal to 1.0."),
    ]

    return pd.DataFrame(rows, columns=["Column", "Description"])


def write_results(
    daily_bwsd: pd.DataFrame,
    exceedance_table: pd.DataFrame,
    parameters_table: pd.DataFrame,
    metadata_table: pd.DataFrame,
    column_dictionary: pd.DataFrame,
    output_xlsx: Path,
) -> None:
    """Write BWSD outputs to an Excel workbook."""
    output_xlsx = Path(output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    daily_to_excel = daily_bwsd.copy()
    daily_to_excel["Date"] = pd.to_datetime(daily_to_excel["Date"]).dt.date

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        daily_to_excel.to_excel(writer, sheet_name="Daily_BWSD", index=False)
        exceedance_table.to_excel(writer, sheet_name="Exceedance", index=False)
        parameters_table.to_excel(writer, sheet_name="Parameters", index=False)
        metadata_table.to_excel(writer, sheet_name="Metadata", index=False)
        column_dictionary.to_excel(writer, sheet_name="Column_Dictionary", index=False)

        for worksheet in writer.sheets.values():
            autofit_worksheet(worksheet)


def print_key_results(metadata_table: pd.DataFrame, output_xlsx: Path) -> None:
    """Print key results to console."""
    lookup = dict(zip(metadata_table["Item"], metadata_table["Value"]))

    print("=" * 88)
    print("BWSD calculation completed")
    print("-" * 88)
    print(f"Output workbook                       : {output_xlsx}")
    print(f"Input dataset DOI                     : {lookup.get('Input_dataset_DOI')}")
    print(f"Date range                            : {lookup.get('Date_range')}")
    print(f"T0                                    : {lookup.get('T0')}")
    print(
        f"Formal evaluation start               : {lookup.get('Formal_evaluation_start')}"
    )
    print(
        f"Initialization max DeltaS             : {lookup.get('Initialization_max_DeltaS_mm')} mm"
    )
    print(f"S_init                                : {lookup.get('S_init_mm')} mm")
    print(f"Max DeltaS                            : {lookup.get('Max_DeltaS_mm')} mm")
    print(f"Date of max DeltaS                    : {lookup.get('Date_of_max_DeltaS')}")
    print(f"Max formal BWSD                       : {lookup.get('Max_BWSD_formal')}")
    print(
        f"Date of max formal BWSD               : {lookup.get('Date_of_max_BWSD_formal')}"
    )
    print(
        f"Formal BWSD >= 0.70 days              : {lookup.get('Formal_BWSD_ge_0.70_days')}"
    )
    print(
        f"Formal BWSD >= 0.90 days              : {lookup.get('Formal_BWSD_ge_0.90_days')}"
    )
    print(
        f"Formal BWSD >= 1.00 days              : {lookup.get('Formal_BWSD_ge_1.00_days')}"
    )
    print("=" * 88)
